"""
BANTMS AP Master Sheet Generator
=================================

Generates the AP_Master (Accounts Payable) sheet with:
- Payment obligation tracking
- Wednesday snap logic (payment day calculation)
- Vendor tier-based snap direction
- Auto-postponement tracking
- Currency conversion to SAR

Wednesday Snap Logic:
- Tier 1 (Critical): Snap to Wednesday BEFORE due date
- Tier 2/3 (Standard/Flexible): Snap to Wednesday AFTER due date
"""

from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import List, Dict, Optional
import random

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.styles import (
    Font, PatternFill, Border, Side, Alignment
)

from ..config.constants import (
    SaudiCalendarConfig,
    CurrencyFramework,
    VendorTierConfig,
    PaymentType,
    PaymentStatus,
    ConfirmationStatus
)
from ..config.styles import (
    ColorPalette, Typography, NumberFormats,
    ColumnWidths, RowHeights
)


# =============================================================================
# SECTION 1: COLUMN DEFINITIONS
# =============================================================================

AP_MASTER_COLUMNS: List[Dict] = [
    {
        "header": "Payment_ID",
        "width": 14,
        "format": "@",
        "alignment": "center",
        "description": "Unique payment identifier",
        "formula": None
    },
    {
        "header": "Vendor_ID",
        "width": 12,
        "format": "@",
        "alignment": "center",
        "description": "Vendor reference code",
        "formula": None
    },
    {
        "header": "Vendor_Name",
        "width": 28,
        "format": "@",
        "alignment": "left",
        "description": "Vendor company name",
        "formula": None
    },
    {
        "header": "Vendor_Tier",
        "width": 10,
        "format": "0",
        "alignment": "center",
        "description": "1=Critical, 2=Standard, 3=Flexible",
        "formula": None
    },
    {
        "header": "Payment_Type",
        "width": 14,
        "format": "@",
        "alignment": "center",
        "description": "Supplier/Government/Strategic/Utilities/Rent",
        "formula": None
    },
    {
        "header": "Invoice_No",
        "width": 16,
        "format": "@",
        "alignment": "left",
        "description": "Vendor invoice number",
        "formula": None
    },
    {
        "header": "Invoice_Date",
        "width": 12,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Invoice date",
        "formula": None
    },
    {
        "header": "Original_Due_Date",
        "width": 14,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Original payment due date",
        "formula": None
    },
    {
        "header": "Snap_Direction",
        "width": 12,
        "format": "@",
        "alignment": "center",
        "description": "BEFORE (Tier 1) or AFTER (Tier 2/3)",
        "formula": '=IF([@Vendor_Tier]=1,"BEFORE","AFTER")'
    },
    {
        "header": "Calculated_Wednesday",
        "width": 16,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Target Wednesday payment date",
        "formula": '=IF(WEEKDAY([@Original_Due_Date],2)=3,[@Original_Due_Date],IF([@Vendor_Tier]=1,[@Original_Due_Date]-MOD(WEEKDAY([@Original_Due_Date],2)-3+7,7),[@Original_Due_Date]+MOD(10-WEEKDAY([@Original_Due_Date],2),7)))'
    },
    {
        "header": "Amount_Original_CCY",
        "width": 16,
        "format": "#,##0.00",
        "alignment": "right",
        "description": "Amount in original currency",
        "formula": None
    },
    {
        "header": "CCY",
        "width": 8,
        "format": "@",
        "alignment": "center",
        "description": "Currency code",
        "formula": None
    },
    {
        "header": "FX_Rate",
        "width": 10,
        "format": "#,##0.0000",
        "alignment": "right",
        "description": "Exchange rate to SAR",
        "formula": '=IFERROR(IF([@CCY]="SAR",1,INDEX(FX_Rate_Lookup,MATCH([@CCY],CCY_Codes,0),3)),1)'
    },
    {
        "header": "Amount_SAR",
        "width": 16,
        "format": "#,##0.00",
        "alignment": "right",
        "description": "Amount in SAR",
        "formula": '=IFERROR([@Amount_Original_CCY]*[@FX_Rate],0)'
    },
    {
        "header": "Status",
        "width": 12,
        "format": "@",
        "alignment": "center",
        "description": "Pending/Paid/Postponed-1/Postponed-2/Urgent",
        "formula": None
    },
    {
        "header": "Postpone_Count",
        "width": 12,
        "format": "0",
        "alignment": "center",
        "description": "Number of times postponed (0-2)",
        "formula": None
    },
    {
        "header": "Last_Status_Change",
        "width": 14,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Date of last status change",
        "formula": None
    },
    {
        "header": "Actual_Payment_Date",
        "width": 16,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Actual payment date (when paid)",
        "formula": None
    },
    {
        "header": "Week_ID",
        "width": 10,
        "format": "00",
        "alignment": "center",
        "description": "Saudi week number for Calculated_Wednesday",
        "formula": '=IFERROR(INT(([@Calculated_Wednesday]-DATE(YEAR([@Calculated_Wednesday]),1,1)-MOD(WEEKDAY(DATE(YEAR([@Calculated_Wednesday]),1,1))+1,7)+1)/7)+1,"")'
    },
    {
        "header": "Days_Until_Due",
        "width": 12,
        "format": "0",
        "alignment": "center",
        "description": "Days until calculated payment date",
        "formula": '=IFERROR([@Calculated_Wednesday]-TODAY(),"")'
    },
    {
        "header": "Priority_Score",
        "width": 12,
        "format": "0.00",
        "alignment": "center",
        "description": "Payment priority score",
        "formula": '=IFERROR((4-[@Vendor_Tier])*100-[@Days_Until_Due]+IF([@Status]="Urgent",500,0),0)'
    },
    {
        "header": "Confirmation",
        "width": 12,
        "format": "@",
        "alignment": "center",
        "description": "Confirmed/Estimated",
        "formula": None
    },
    {
        "header": "Notes",
        "width": 30,
        "format": "@",
        "alignment": "left",
        "description": "Additional notes",
        "formula": None
    },
]


# =============================================================================
# SECTION 2: SAMPLE DATA GENERATOR
# =============================================================================

def generate_sample_ap_data() -> List[Dict]:
    """
    Generate sample AP data for demonstration.

    Returns:
        List of payment dictionaries with realistic Saudi business data.
    """
    # Sample vendors by tier
    vendors = {
        1: [  # Tier 1 - Critical
            {"id": "V-001", "name": "Saudi Aramco", "type": "Strategic"},
            {"id": "V-002", "name": "SABIC", "type": "Strategic"},
            {"id": "V-003", "name": "Saudi Electricity Company", "type": "Utilities"},
            {"id": "V-004", "name": "GOSI", "type": "Government"},
            {"id": "V-005", "name": "ZATCA", "type": "Government"},
        ],
        2: [  # Tier 2 - Standard
            {"id": "V-010", "name": "Al-Marai Company", "type": "Supplier"},
            {"id": "V-011", "name": "Jarir Marketing", "type": "Supplier"},
            {"id": "V-012", "name": "Extra Stores", "type": "Supplier"},
            {"id": "V-013", "name": "Panda Retail", "type": "Supplier"},
            {"id": "V-014", "name": "STC Telecommunications", "type": "Utilities"},
        ],
        3: [  # Tier 3 - Flexible
            {"id": "V-020", "name": "Office Supplies Co.", "type": "Supplier"},
            {"id": "V-021", "name": "Building Maintenance Ltd", "type": "Supplier"},
            {"id": "V-022", "name": "IT Solutions Provider", "type": "Supplier"},
            {"id": "V-023", "name": "Cleaning Services Co.", "type": "Supplier"},
            {"id": "V-024", "name": "Security Services LLC", "type": "Supplier"},
        ],
    }

    payments = []
    base_date = date.today() - timedelta(days=30)
    payment_id = 1

    # Generate payments across tiers
    for tier, tier_vendors in vendors.items():
        for vendor in tier_vendors:
            # Generate 2-4 payments per vendor
            num_payments = random.randint(2, 4)

            for i in range(num_payments):
                invoice_date = base_date + timedelta(days=random.randint(0, 60))
                due_date = invoice_date + timedelta(days=random.choice([30, 45, 60]))

                # Amount varies by tier
                if tier == 1:
                    amount = random.randint(200000, 1000000)
                elif tier == 2:
                    amount = random.randint(50000, 300000)
                else:
                    amount = random.randint(5000, 80000)

                # Currency (mostly SAR, some USD for tier 1)
                if tier == 1 and random.random() < 0.3:
                    ccy = "USD"
                else:
                    ccy = "SAR"

                # Status distribution
                status_options = ["Pending"] * 6 + ["Paid"] * 2 + ["Postponed-1"] + ["Urgent"]
                status = random.choice(status_options)

                postpone_count = 0
                if "Postponed" in status:
                    postpone_count = int(status[-1])

                actual_payment_date = None
                if status == "Paid":
                    actual_payment_date = due_date + timedelta(days=random.randint(-3, 3))

                payments.append({
                    "Payment_ID": f"AP-{payment_id:06d}",
                    "Vendor_ID": vendor["id"],
                    "Vendor_Name": vendor["name"],
                    "Vendor_Tier": tier,
                    "Payment_Type": vendor["type"],
                    "Invoice_No": f"INV-{invoice_date.strftime('%Y%m')}-{payment_id:04d}",
                    "Invoice_Date": invoice_date,
                    "Original_Due_Date": due_date,
                    "Amount_Original_CCY": float(amount),
                    "CCY": ccy,
                    "Status": status,
                    "Postpone_Count": postpone_count,
                    "Last_Status_Change": date.today() - timedelta(days=random.randint(0, 10)),
                    "Actual_Payment_Date": actual_payment_date,
                    "Confirmation": random.choice(["Confirmed", "Estimated"]),
                    "Notes": "",
                })
                payment_id += 1

    return payments


# =============================================================================
# SECTION 3: SHEET GENERATOR
# =============================================================================

class APMasterSheetGenerator:
    """
    Generator for AP_Master (Accounts Payable) sheet.

    Creates a formatted Excel sheet with:
    - Payment obligation tracking
    - Wednesday snap logic formulas
    - Vendor tier-based payment scheduling
    - Currency conversion to SAR
    - Auto-calculation of priority scores
    """

    SHEET_NAME = "AP_Master"
    TABLE_NAME = "tbl_AP_Master"
    TITLE_ROW = 1
    HEADER_ROW = 3
    DATA_START_ROW = 4

    def __init__(self, workbook: Workbook):
        """
        Initialize the AP Master sheet generator.

        Args:
            workbook: The openpyxl Workbook to add the sheet to.
        """
        self.workbook = workbook
        self.sheet: Optional[Worksheet] = None
        self.payments: List[Dict] = []
        self.last_data_row = self.DATA_START_ROW

    def generate(
        self,
        payments: Optional[List[Dict]] = None,
        include_sample_data: bool = True
    ) -> Worksheet:
        """
        Generate the AP_Master sheet.

        Args:
            payments: Optional list of payment dictionaries.
            include_sample_data: Whether to include sample data (default True).

        Returns:
            The generated Worksheet.
        """
        if payments:
            self.payments = payments
        elif include_sample_data:
            self.payments = generate_sample_ap_data()
        else:
            self.payments = []

        # Create or get sheet
        if self.SHEET_NAME in self.workbook.sheetnames:
            self.sheet = self.workbook[self.SHEET_NAME]
        else:
            self.sheet = self.workbook.create_sheet(self.SHEET_NAME)

        # Build sheet components
        self._write_title()
        self._write_headers()
        self._write_data()
        self._apply_conditional_formatting()
        self._create_table()
        self._set_column_widths()
        self._add_data_validation()
        self._freeze_panes()
        self._add_instructions()

        return self.sheet

    def _write_title(self):
        """Write sheet title and description."""
        self.sheet.cell(
            row=self.TITLE_ROW,
            column=1,
            value="AP MASTER - Accounts Payable Management"
        )
        title_cell = self.sheet.cell(row=self.TITLE_ROW, column=1)
        title_cell.font = Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_TITLE,
            bold=True,
            color=ColorPalette.HEADER_DARK_NAVY
        )

        # Merge title cells
        self.sheet.merge_cells(
            start_row=self.TITLE_ROW,
            start_column=1,
            end_row=self.TITLE_ROW,
            end_column=len(AP_MASTER_COLUMNS)
        )

        # Subtitle
        self.sheet.cell(
            row=self.TITLE_ROW + 1,
            column=1,
            value="Wednesday Payment Scheduling | Vendor Tier Priority | All amounts in SAR"
        )
        subtitle_cell = self.sheet.cell(row=self.TITLE_ROW + 1, column=1)
        subtitle_cell.font = Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_SMALL,
            italic=True,
            color=ColorPalette.TEXT_DARK_GRAY
        )

    def _write_headers(self):
        """Write column headers."""
        header_style = {
            "font": Font(
                name=Typography.FONT_FAMILY,
                size=Typography.SIZE_HEADER,
                bold=True,
                color=ColorPalette.TEXT_WHITE
            ),
            "fill": PatternFill(
                start_color=ColorPalette.HEADER_DARK_NAVY,
                end_color=ColorPalette.HEADER_DARK_NAVY,
                fill_type="solid"
            ),
            "alignment": Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            ),
            "border": Border(
                bottom=Side(style="medium", color=ColorPalette.BORDER_MEDIUM_BLACK)
            )
        }

        for col_idx, col_def in enumerate(AP_MASTER_COLUMNS, start=1):
            cell = self.sheet.cell(
                row=self.HEADER_ROW,
                column=col_idx,
                value=col_def["header"]
            )

            cell.font = header_style["font"]
            cell.fill = header_style["fill"]
            cell.alignment = header_style["alignment"]
            cell.border = header_style["border"]

    def _write_data(self):
        """Write payment data rows with formulas."""
        for row_idx, payment in enumerate(self.payments):
            excel_row = self.DATA_START_ROW + row_idx
            is_alternate = row_idx % 2 == 1

            self._write_payment_row(excel_row, payment, is_alternate)

        self.last_data_row = self.DATA_START_ROW + len(self.payments) - 1
        if self.last_data_row < self.DATA_START_ROW:
            self.last_data_row = self.DATA_START_ROW

    def _write_payment_row(self, row: int, payment: Dict, is_alternate: bool):
        """Write a single payment row."""
        fill_color = ColorPalette.ROW_LIGHT_GRAY if is_alternate else ColorPalette.ROW_WHITE

        # Determine status-based styling
        status = payment.get("Status", "Pending")
        status_colors = {
            "Paid": ColorPalette.STATUS_PAID_GREEN,
            "Pending": fill_color,
            "Postponed-1": ColorPalette.STATUS_POSTPONED_1_ORANGE,
            "Postponed-2": ColorPalette.STATUS_POSTPONED_2_RED,
            "Urgent": ColorPalette.STATUS_URGENT_PURPLE,
        }
        row_fill_color = status_colors.get(status, fill_color)

        style = {
            "font": Font(
                name=Typography.FONT_FAMILY,
                size=Typography.SIZE_BODY,
                color=ColorPalette.TEXT_BLACK
            ),
            "fill": PatternFill(
                start_color=row_fill_color,
                end_color=row_fill_color,
                fill_type="solid"
            ),
            "border": Border(
                left=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
                right=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
                top=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
                bottom=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY)
            )
        }

        for col_idx, col_def in enumerate(AP_MASTER_COLUMNS, start=1):
            header = col_def["header"]

            # Determine cell value
            if col_def["formula"]:
                value = col_def["formula"]
            elif header in payment:
                value = payment[header]
            else:
                value = None

            cell = self.sheet.cell(row=row, column=col_idx, value=value)
            cell.font = style["font"]
            cell.fill = style["fill"]
            cell.border = style["border"]

            # Apply alignment
            h_align = col_def["alignment"]
            cell.alignment = Alignment(horizontal=h_align, vertical="center")

            # Apply number format
            cell.number_format = col_def["format"]

    def _apply_conditional_formatting(self):
        """Apply conditional formatting rules."""
        # Status-based row coloring is handled in _write_payment_row
        # Additional conditional formatting can be added here
        pass

    def _create_table(self):
        """Create Excel table for AP data."""
        start_col = "A"
        end_col = get_column_letter(len(AP_MASTER_COLUMNS))

        table_end_row = max(self.last_data_row, self.DATA_START_ROW)
        table_range = f"{start_col}{self.HEADER_ROW}:{end_col}{table_end_row}"

        table = Table(
            displayName=self.TABLE_NAME,
            ref=table_range
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium3",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        table.tableStyleInfo = table_style

        self.sheet.add_table(table)

    def _set_column_widths(self):
        """Set column widths based on definitions."""
        for col_idx, col_def in enumerate(AP_MASTER_COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            self.sheet.column_dimensions[col_letter].width = col_def["width"]

        # Set row heights
        self.sheet.row_dimensions[self.TITLE_ROW].height = RowHeights.TITLE
        self.sheet.row_dimensions[self.HEADER_ROW].height = RowHeights.HEADER

    def _add_data_validation(self):
        """Add data validation for input columns."""
        # Vendor Tier dropdown
        tier_validation = DataValidation(
            type="list",
            formula1='"1,2,3"',
            allow_blank=False
        )
        tier_col = get_column_letter(4)
        tier_validation.add(f"{tier_col}{self.DATA_START_ROW}:{tier_col}5000")
        self.sheet.add_data_validation(tier_validation)

        # Payment Type dropdown
        payment_types = ",".join([pt.value for pt in PaymentType])
        type_validation = DataValidation(
            type="list",
            formula1=f'"{payment_types}"',
            allow_blank=False
        )
        type_col = get_column_letter(5)
        type_validation.add(f"{type_col}{self.DATA_START_ROW}:{type_col}5000")
        self.sheet.add_data_validation(type_validation)

        # Currency dropdown
        ccy_codes = ",".join(CurrencyFramework.CURRENCY_CODES)
        ccy_validation = DataValidation(
            type="list",
            formula1=f'"{ccy_codes}"',
            allow_blank=False
        )
        ccy_col = get_column_letter(12)
        ccy_validation.add(f"{ccy_col}{self.DATA_START_ROW}:{ccy_col}5000")
        self.sheet.add_data_validation(ccy_validation)

        # Status dropdown
        status_options = ",".join([s.value for s in PaymentStatus])
        status_validation = DataValidation(
            type="list",
            formula1=f'"{status_options}"',
            allow_blank=False
        )
        status_col = get_column_letter(15)
        status_validation.add(f"{status_col}{self.DATA_START_ROW}:{status_col}5000")
        self.sheet.add_data_validation(status_validation)

        # Confirmation dropdown
        confirm_options = ",".join([c.value for c in ConfirmationStatus])
        confirm_validation = DataValidation(
            type="list",
            formula1=f'"{confirm_options}"',
            allow_blank=False
        )
        confirm_col = get_column_letter(22)
        confirm_validation.add(f"{confirm_col}{self.DATA_START_ROW}:{confirm_col}5000")
        self.sheet.add_data_validation(confirm_validation)

    def _freeze_panes(self):
        """Freeze header rows and key columns."""
        self.sheet.freeze_panes = f"D{self.DATA_START_ROW}"

    def _add_instructions(self):
        """Add usage instructions at the bottom."""
        instructions_row = self.last_data_row + 3

        instructions = [
            "",
            "WEDNESDAY SNAP LOGIC:",
            "• All payments scheduled for Wednesday (Saudi payment day)",
            "• Tier 1 (Critical): Snap to Wednesday BEFORE due date",
            "• Tier 2/3 (Standard/Flexible): Snap to Wednesday AFTER due date",
            "• Formula: If due date IS Wednesday, use as-is",
            "",
            "VENDOR TIERS:",
            "• Tier 1: Mission-critical (Aramco, SABIC, Government)",
            "• Tier 2: Standard suppliers (regular business partners)",
            "• Tier 3: Flexible (office supplies, maintenance)",
            "",
            "STATUS CODES:",
            "• Pending: Awaiting payment",
            "• Paid: Payment completed",
            "• Postponed-1: First postponement",
            "• Postponed-2: Second postponement (max)",
            "• Urgent: Priority payment required",
            "",
            "PRIORITY SCORE:",
            "• Higher score = higher priority",
            "• Based on: Tier, Days until due, Urgent flag",
        ]

        for idx, text in enumerate(instructions):
            cell = self.sheet.cell(row=instructions_row + idx, column=1, value=text)
            if text.startswith("WEDNESDAY") or text.startswith("VENDOR") or text.startswith("STATUS") or text.startswith("PRIORITY"):
                cell.font = Font(
                    name=Typography.FONT_FAMILY,
                    size=Typography.SIZE_BODY,
                    bold=True,
                    color=ColorPalette.HEADER_DARK_NAVY
                )
            else:
                cell.font = Font(
                    name=Typography.FONT_FAMILY,
                    size=Typography.SIZE_SMALL,
                    color=ColorPalette.TEXT_DARK_GRAY
                )


# =============================================================================
# SECTION 4: UTILITY FUNCTIONS
# =============================================================================

def generate_ap_master_sheet(
    workbook: Workbook,
    include_sample_data: bool = True
) -> Worksheet:
    """
    Convenience function to generate AP_Master sheet.

    Args:
        workbook: The openpyxl Workbook.
        include_sample_data: Whether to include sample data.

    Returns:
        The generated Worksheet.
    """
    generator = APMasterSheetGenerator(workbook)
    return generator.generate(include_sample_data=include_sample_data)


def get_wednesday_snap_formula() -> str:
    """
    Get the Wednesday snap formula for payment date calculation.

    Logic:
    - If due date IS Wednesday: Use as-is
    - If Tier 1: Snap to Wednesday BEFORE
    - If Tier 2/3: Snap to Wednesday AFTER

    Returns:
        Excel formula string for Wednesday snap calculation.
    """
    return (
        '=IF(WEEKDAY([@Original_Due_Date],2)=3,'
        '[@Original_Due_Date],'
        'IF([@Vendor_Tier]=1,'
        '[@Original_Due_Date]-MOD(WEEKDAY([@Original_Due_Date],2)-3+7,7),'
        '[@Original_Due_Date]+MOD(10-WEEKDAY([@Original_Due_Date],2),7)))'
    )


def get_snap_direction_formula() -> str:
    """
    Get the snap direction formula based on vendor tier.

    Returns:
        Excel formula string for snap direction.
    """
    return '=IF([@Vendor_Tier]=1,"BEFORE","AFTER")'


def get_priority_score_formula() -> str:
    """
    Get the priority score calculation formula.

    Higher score = higher priority based on:
    - Vendor tier (Tier 1 = highest)
    - Days until due (closer = higher)
    - Urgent flag (+500 bonus)

    Returns:
        Excel formula string for priority score.
    """
    return '=IFERROR((4-[@Vendor_Tier])*100-[@Days_Until_Due]+IF([@Status]="Urgent",500,0),0)'
