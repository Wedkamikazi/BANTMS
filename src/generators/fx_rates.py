"""
BANTMS FX Rates Sheet Generator
================================

Generates the FX_Rates reference sheet with:
- Currency codes and names
- Exchange rates to SAR (base currency)
- Effective dates
- Data validation and formatting

All final calculations throughout the workbook convert to SAR.
"""

from datetime import date, datetime
from decimal import Decimal
from typing import List, Dict, Optional, Tuple

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.styles import (
    Font, Fill, PatternFill, Border, Side, Alignment, Protection, NamedStyle
)
from openpyxl.formatting.rule import FormulaRule

from ..config.constants import CurrencyFramework, CurrencyInfo
from ..config.styles import (
    ColorPalette, Typography, NumberFormats,
    ColumnWidths, RowHeights, CellStyles, Alignments
)


# =============================================================================
# SECTION 1: FX RATES DATA STRUCTURE
# =============================================================================

class FXRateRecord:
    """
    Represents a single FX rate record.

    Attributes:
        ccy_code: Currency code (e.g., "USD")
        ccy_name: Currency full name
        rate_to_sar: Exchange rate to SAR
        effective_date: Date when rate became effective
        source: Source of the rate (e.g., "SAMA", "Reuters")
        bid_rate: Optional bid rate for spread calculation
        ask_rate: Optional ask rate for spread calculation
    """

    def __init__(
        self,
        ccy_code: str,
        ccy_name: str,
        rate_to_sar: Decimal,
        effective_date: date,
        source: str = "Manual Entry",
        bid_rate: Optional[Decimal] = None,
        ask_rate: Optional[Decimal] = None
    ):
        self.ccy_code = ccy_code
        self.ccy_name = ccy_name
        self.rate_to_sar = rate_to_sar
        self.effective_date = effective_date
        self.source = source
        self.bid_rate = bid_rate or rate_to_sar
        self.ask_rate = ask_rate or rate_to_sar


# =============================================================================
# SECTION 2: COLUMN DEFINITIONS
# =============================================================================

FX_RATES_COLUMNS: List[Dict] = [
    {
        "header": "CCY_Code",
        "width": 10,
        "format": "@",
        "alignment": "center",
        "description": "ISO Currency Code"
    },
    {
        "header": "CCY_Name",
        "width": 25,
        "format": "@",
        "alignment": "left",
        "description": "Currency Full Name"
    },
    {
        "header": "Rate_to_SAR",
        "width": 15,
        "format": "#,##0.0000",
        "alignment": "right",
        "description": "Exchange Rate to SAR"
    },
    {
        "header": "Effective_Date",
        "width": 14,
        "format": "DD-MMM-YYYY",
        "alignment": "center",
        "description": "Rate Effective Date"
    },
    {
        "header": "Source",
        "width": 18,
        "format": "@",
        "alignment": "left",
        "description": "Rate Source"
    },
    {
        "header": "Bid_Rate",
        "width": 12,
        "format": "#,##0.0000",
        "alignment": "right",
        "description": "Bid Rate (for spread)"
    },
    {
        "header": "Ask_Rate",
        "width": 12,
        "format": "#,##0.0000",
        "alignment": "right",
        "description": "Ask Rate (for spread)"
    },
    {
        "header": "Spread_BPS",
        "width": 12,
        "format": "#,##0.00",
        "alignment": "right",
        "description": "Spread in Basis Points"
    },
    {
        "header": "Last_Updated",
        "width": 18,
        "format": "DD-MMM-YYYY HH:MM",
        "alignment": "center",
        "description": "Last Update Timestamp"
    },
]


# =============================================================================
# SECTION 3: DEFAULT FX RATES DATA
# =============================================================================

def get_default_fx_rates() -> List[FXRateRecord]:
    """
    Get default FX rates based on CurrencyFramework configuration.

    Returns:
        List of FXRateRecord objects with default rates.

    Note:
        SAR is the base currency with rate 1.0000.
        All other currencies are converted TO SAR.
    """
    today = date.today()
    rates = []

    for ccy_code, ccy_info in CurrencyFramework.CURRENCIES.items():
        # Calculate bid/ask with small spread for non-SAR currencies
        if ccy_code == "SAR":
            bid = ask = Decimal("1.0000")
        else:
            base_rate = ccy_info.rate_to_sar
            spread = base_rate * Decimal("0.001")  # 0.1% spread
            bid = base_rate - spread / 2
            ask = base_rate + spread / 2

        rates.append(FXRateRecord(
            ccy_code=ccy_info.code,
            ccy_name=ccy_info.name,
            rate_to_sar=ccy_info.rate_to_sar,
            effective_date=today,
            source="System Default",
            bid_rate=bid,
            ask_rate=ask
        ))

    return rates


# =============================================================================
# SECTION 4: STYLE DEFINITIONS
# =============================================================================

def create_header_style() -> Dict:
    """Create header row style configuration."""
    return {
        "font": Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_HEADER,
            bold=True,
            color=ColorPalette.TEXT_WHITE
        ),
        "fill": PatternFill(
            start_color=ColorPalette.HEADER_DARK_NAVY,
            end_color=ColorPalette.HEADER_DARK_NAVY,
            fill_type="solid"
        ),
        "alignment": Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        ),
        "border": Border(
            bottom=Side(style="medium", color=ColorPalette.BORDER_MEDIUM_BLACK)
        )
    }


def create_data_style(is_alternate: bool = False) -> Dict:
    """Create data row style configuration."""
    fill_color = ColorPalette.ROW_LIGHT_GRAY if is_alternate else ColorPalette.ROW_WHITE
    return {
        "font": Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_BODY,
            color=ColorPalette.TEXT_BLACK
        ),
        "fill": PatternFill(
            start_color=fill_color,
            end_color=fill_color,
            fill_type="solid"
        ),
        "border": Border(
            left=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
            right=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
            top=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
            bottom=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY)
        )
    }


def create_sar_highlight_style() -> Dict:
    """Create style for SAR row (base currency highlight)."""
    return {
        "font": Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_BODY,
            bold=True,
            color=ColorPalette.TEXT_BLACK
        ),
        "fill": PatternFill(
            start_color=ColorPalette.HEADER_LIGHT_BLUE,
            end_color=ColorPalette.HEADER_LIGHT_BLUE,
            fill_type="solid"
        ),
        "border": Border(
            left=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
            right=Side(style="thin", color=ColorPalette.BORDER_THIN_GRAY),
            top=Side(style="medium", color=ColorPalette.HEADER_MEDIUM_BLUE),
            bottom=Side(style="medium", color=ColorPalette.HEADER_MEDIUM_BLUE)
        )
    }


# =============================================================================
# SECTION 5: FX RATES SHEET GENERATOR
# =============================================================================

class FXRatesSheetGenerator:
    """
    Generator for FX_Rates reference sheet.

    Creates a formatted Excel sheet with:
    - Currency reference data
    - Exchange rates to SAR
    - Bid/Ask spreads
    - Data validation
    - Named ranges for formula references
    """

    SHEET_NAME = "FX_Rates"
    TABLE_NAME = "tbl_FX_Rates"
    TITLE_ROW = 1
    HEADER_ROW = 3
    DATA_START_ROW = 4

    def __init__(self, workbook: Workbook):
        """
        Initialize the FX Rates sheet generator.

        Args:
            workbook: The openpyxl Workbook to add the sheet to.
        """
        self.workbook = workbook
        self.sheet: Optional[Worksheet] = None
        self.rates: List[FXRateRecord] = []

    def generate(self, rates: Optional[List[FXRateRecord]] = None) -> Worksheet:
        """
        Generate the FX_Rates sheet.

        Args:
            rates: Optional list of FX rate records. Uses defaults if not provided.

        Returns:
            The generated Worksheet.
        """
        self.rates = rates or get_default_fx_rates()

        # Create or get sheet
        if self.SHEET_NAME in self.workbook.sheetnames:
            self.sheet = self.workbook[self.SHEET_NAME]
        else:
            self.sheet = self.workbook.create_sheet(self.SHEET_NAME)

        # Build sheet components
        self._write_title()
        self._write_headers()
        self._write_data()
        self._apply_formatting()
        self._create_table()
        self._set_column_widths()
        self._create_named_ranges()
        self._add_data_validation()
        self._add_instructions()

        return self.sheet

    def _write_title(self):
        """Write sheet title and description."""
        # Title
        self.sheet.cell(row=self.TITLE_ROW, column=1, value="FX RATES - Currency Exchange Reference")
        title_cell = self.sheet.cell(row=self.TITLE_ROW, column=1)
        title_cell.font = Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_TITLE,
            bold=True,
            color=ColorPalette.HEADER_DARK_NAVY
        )

        # Merge title cells
        self.sheet.merge_cells(
            start_row=self.TITLE_ROW,
            start_column=1,
            end_row=self.TITLE_ROW,
            end_column=len(FX_RATES_COLUMNS)
        )

        # Subtitle with base currency note
        self.sheet.cell(
            row=self.TITLE_ROW + 1,
            column=1,
            value="Base Currency: SAR (Saudi Riyal) | All amounts convert to SAR"
        )
        subtitle_cell = self.sheet.cell(row=self.TITLE_ROW + 1, column=1)
        subtitle_cell.font = Font(
            name=Typography.FONT_FAMILY,
            size=Typography.SIZE_SMALL,
            italic=True,
            color=ColorPalette.TEXT_DARK_GRAY
        )

    def _write_headers(self):
        """Write column headers."""
        header_style = create_header_style()

        for col_idx, col_def in enumerate(FX_RATES_COLUMNS, start=1):
            cell = self.sheet.cell(
                row=self.HEADER_ROW,
                column=col_idx,
                value=col_def["header"]
            )

            # Apply header styling
            cell.font = header_style["font"]
            cell.fill = header_style["fill"]
            cell.alignment = header_style["alignment"]
            cell.border = header_style["border"]

    def _write_data(self):
        """Write FX rate data rows."""
        now = datetime.now()

        for row_idx, rate in enumerate(self.rates):
            excel_row = self.DATA_START_ROW + row_idx
            is_alternate = row_idx % 2 == 1
            is_sar = rate.ccy_code == "SAR"

            # Get appropriate style
            if is_sar:
                style = create_sar_highlight_style()
            else:
                style = create_data_style(is_alternate)

            # Column A: CCY_Code
            self._write_cell(excel_row, 1, rate.ccy_code, style, "center")

            # Column B: CCY_Name
            self._write_cell(excel_row, 2, rate.ccy_name, style, "left")

            # Column C: Rate_to_SAR
            self._write_cell(excel_row, 3, float(rate.rate_to_sar), style, "right")

            # Column D: Effective_Date
            self._write_cell(excel_row, 4, rate.effective_date, style, "center")

            # Column E: Source
            self._write_cell(excel_row, 5, rate.source, style, "left")

            # Column F: Bid_Rate
            self._write_cell(excel_row, 6, float(rate.bid_rate), style, "right")

            # Column G: Ask_Rate
            self._write_cell(excel_row, 7, float(rate.ask_rate), style, "right")

            # Column H: Spread_BPS (formula)
            spread_formula = f"=IF(F{excel_row}=0,0,(G{excel_row}-F{excel_row})/F{excel_row}*10000)"
            self._write_cell(excel_row, 8, spread_formula, style, "right")

            # Column I: Last_Updated
            self._write_cell(excel_row, 9, now, style, "center")

    def _write_cell(
        self,
        row: int,
        col: int,
        value,
        style: Dict,
        h_align: str
    ):
        """Write a single cell with value and style."""
        cell = self.sheet.cell(row=row, column=col, value=value)
        cell.font = style["font"]
        cell.fill = style["fill"]
        cell.border = style["border"]
        cell.alignment = Alignment(horizontal=h_align, vertical="center")

    def _apply_formatting(self):
        """Apply number formats to data columns."""
        for row_idx in range(len(self.rates)):
            excel_row = self.DATA_START_ROW + row_idx

            for col_idx, col_def in enumerate(FX_RATES_COLUMNS, start=1):
                cell = self.sheet.cell(row=excel_row, column=col_idx)
                cell.number_format = col_def["format"]

    def _create_table(self):
        """Create Excel table for the FX rates data."""
        if len(self.rates) == 0:
            return

        # Define table range
        start_col = "A"
        end_col = get_column_letter(len(FX_RATES_COLUMNS))
        end_row = self.DATA_START_ROW + len(self.rates) - 1

        table_range = f"{start_col}{self.HEADER_ROW}:{end_col}{end_row}"

        # Create table
        table = Table(
            displayName=self.TABLE_NAME,
            ref=table_range
        )

        # Apply table style
        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        table.tableStyleInfo = table_style

        self.sheet.add_table(table)

    def _set_column_widths(self):
        """Set column widths based on definitions."""
        for col_idx, col_def in enumerate(FX_RATES_COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            self.sheet.column_dimensions[col_letter].width = col_def["width"]

        # Set row heights
        self.sheet.row_dimensions[self.TITLE_ROW].height = RowHeights.TITLE
        self.sheet.row_dimensions[self.HEADER_ROW].height = RowHeights.HEADER

    def _create_named_ranges(self):
        """Create named ranges for formula references throughout workbook."""
        # Named range for currency codes (for data validation)
        ccy_range = f"'{self.SHEET_NAME}'!$A${self.DATA_START_ROW}:$A${self.DATA_START_ROW + len(self.rates) - 1}"
        self.workbook.defined_names.add(
            self.workbook.defined_names.new(name="CCY_Codes", value=ccy_range)
        )

        # Named range for rates lookup
        rates_range = f"'{self.SHEET_NAME}'!$A${self.DATA_START_ROW}:$C${self.DATA_START_ROW + len(self.rates) - 1}"
        self.workbook.defined_names.add(
            self.workbook.defined_names.new(name="FX_Rate_Lookup", value=rates_range)
        )

    def _add_data_validation(self):
        """Add data validation for source column."""
        # Source dropdown
        source_validation = DataValidation(
            type="list",
            formula1='"System Default,SAMA,Reuters,Bloomberg,Bank Rate,Manual Entry"',
            allow_blank=True
        )
        source_validation.error = "Please select a valid source"
        source_validation.errorTitle = "Invalid Source"

        # Apply to source column
        source_col = get_column_letter(5)  # Column E
        source_range = f"{source_col}{self.DATA_START_ROW}:{source_col}{self.DATA_START_ROW + 100}"
        source_validation.add(source_range)
        self.sheet.add_data_validation(source_validation)

    def _add_instructions(self):
        """Add usage instructions at the bottom of the sheet."""
        instructions_row = self.DATA_START_ROW + len(self.rates) + 2

        instructions = [
            "",
            "USAGE INSTRUCTIONS:",
            "• SAR is the base currency - all amounts in other sheets convert to SAR",
            "• Rate_to_SAR: Multiply foreign currency amount by this rate to get SAR equivalent",
            "• Example: 1,000 USD × 3.7500 = 3,750 SAR",
            "• Update Effective_Date when changing rates",
            "• Spread_BPS shows bid-ask spread in basis points",
            "",
            "FORMULA FOR CURRENCY CONVERSION (use in other sheets):",
            "=IFERROR(IF([@CCY]=\"SAR\",[@Amount],[@Amount]*INDEX(FX_Rate_Lookup,MATCH([@CCY],CCY_Codes,0),3)),0)",
        ]

        for idx, text in enumerate(instructions):
            cell = self.sheet.cell(row=instructions_row + idx, column=1, value=text)
            if text.startswith("USAGE") or text.startswith("FORMULA"):
                cell.font = Font(
                    name=Typography.FONT_FAMILY,
                    size=Typography.SIZE_BODY,
                    bold=True,
                    color=ColorPalette.HEADER_DARK_NAVY
                )
            else:
                cell.font = Font(
                    name=Typography.FONT_FAMILY,
                    size=Typography.SIZE_SMALL,
                    color=ColorPalette.TEXT_DARK_GRAY
                )


# =============================================================================
# SECTION 6: UTILITY FUNCTIONS
# =============================================================================

def generate_fx_rates_sheet(workbook: Workbook) -> Worksheet:
    """
    Convenience function to generate FX_Rates sheet.

    Args:
        workbook: The openpyxl Workbook.

    Returns:
        The generated Worksheet.
    """
    generator = FXRatesSheetGenerator(workbook)
    return generator.generate()


def get_fx_conversion_formula(
    amount_column: str,
    ccy_column: str,
    row_ref: str = "@"
) -> str:
    """
    Generate the FX conversion formula for use in other sheets.

    Args:
        amount_column: Column containing the original amount
        ccy_column: Column containing the currency code
        row_ref: Row reference style ("@" for structured refs, or row number)

    Returns:
        Excel formula string for currency conversion to SAR.
    """
    if row_ref == "@":
        # Structured table reference
        return (
            f'=IFERROR('
            f'IF([{ccy_column}]="SAR",'
            f'[{amount_column}],'
            f'[{amount_column}]*INDEX(FX_Rate_Lookup,MATCH([{ccy_column}],CCY_Codes,0),3)),'
            f'0)'
        )
    else:
        # Cell reference
        return (
            f'=IFERROR('
            f'IF({ccy_column}{row_ref}="SAR",'
            f'{amount_column}{row_ref},'
            f'{amount_column}{row_ref}*INDEX(FX_Rate_Lookup,MATCH({ccy_column}{row_ref},CCY_Codes,0),3)),'
            f'0)'
        )
